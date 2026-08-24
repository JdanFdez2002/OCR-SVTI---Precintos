from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any


@dataclass
class RegistroEsperado:
    patente: str = ""
    marca: str = ""
    fecha_esperada: date | None = None
    precinto_esperado: str = ""
    estado: str = "Pendiente"


_ALIAS_COLUMNAS = {
    "patente": {
        "patente",
        "placa",
        "ppu",
        "patente camion",
        "patente camión",
    },
    "marca": {
        "marca",
        "marca camion",
        "marca camión",
        "camion",
        "camión",
    },
    "fecha_esperada": {
        "fecha",
        "fecha llegada",
        "fecha esperada",
        "fecha llegada esperada",
        "llegada esperada",
    },
    "precinto_esperado": {
        "precinto",
        "codigo precinto",
        "código precinto",
        "precinto esperado",
        "codigo de precinto esperado",
        "código de precinto esperado",
    },
}


def _normalizar(valor: Any) -> str:
    texto = str(valor or "").strip().lower()
    reemplazos = str.maketrans("áéíóúüñ", "aeiouun")
    return " ".join(texto.translate(reemplazos).split())


def _texto(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _fecha(valor: Any) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _texto(valor)
    if not texto:
        return None
    for formato in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            pass
    return None


def _mapear_columnas(encabezados: list[Any]) -> dict[str, int]:
    normalizados = [_normalizar(celda) for celda in encabezados]
    columnas: dict[str, int] = {}
    for campo, aliases in _ALIAS_COLUMNAS.items():
        aliases_norm = {_normalizar(alias) for alias in aliases}
        for idx, encabezado in enumerate(normalizados):
            if encabezado in aliases_norm:
                columnas[campo] = idx
                break
    return columnas


def cargar_excel(ruta: str | Path) -> list[RegistroEsperado]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Falta instalar openpyxl para leer archivos Excel.") from exc

    wb = load_workbook(ruta, read_only=True, data_only=True)
    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []

    columnas = _mapear_columnas(list(filas[0]))
    requeridas = {"patente", "fecha_esperada", "precinto_esperado"}
    faltantes = sorted(requeridas - set(columnas))
    if faltantes:
        raise ValueError(
            "Faltan columnas requeridas: " + ", ".join(faltantes)
        )

    registros: list[RegistroEsperado] = []
    hoy = date.today()
    for fila in filas[1:]:
        patente = _texto(fila[columnas["patente"]]).upper()
        marca = _texto(fila[columnas["marca"]]) if "marca" in columnas else ""
        fecha_esperada = _fecha(fila[columnas["fecha_esperada"]])
        precinto = _texto(fila[columnas["precinto_esperado"]]).upper()
        if not any((patente, marca, fecha_esperada, precinto)):
            continue
        estado = "Atrasado" if fecha_esperada and fecha_esperada < hoy else "Pendiente"
        registros.append(
            RegistroEsperado(
                patente=patente,
                marca=marca,
                fecha_esperada=fecha_esperada,
                precinto_esperado=precinto,
                estado=estado,
            )
        )
    return registros
