import cv2
import numpy as np
import os
import threading
import json
import re
import openpyxl
import glob
from datetime import datetime
from config import *
from modulos.camera_service import CameraService
from modulos.ocr_engine import mejor_ocr

# --- CONFIGURACIÓN DE RUTAS ---
CARPETA_MANIFIESTOS = "manifiestos_diarios"

# ==============================================================================
# FUNCIONES DE MANEJO DE EXCEL (DINÁMICAS)
# ==============================================================================

def obtener_ruta_manifiesto():
    """Retorna la ruta del archivo Excel de hoy dentro de la carpeta dedicada."""
    if not os.path.exists(CARPETA_MANIFIESTOS):
        os.makedirs(CARPETA_MANIFIESTOS)
        
    fecha_hoy = datetime.now().strftime("%Y%m%d")
    nombre_esperado = f"manifiesto_{fecha_hoy}.xlsx"
    ruta_manifiesto = os.path.join(CARPETA_MANIFIESTOS, nombre_esperado)
    
    # Backup: Si no está el de hoy, usar el más reciente
    if not os.path.exists(ruta_manifiesto):
        archivos_excel = glob.glob(os.path.join(CARPETA_MANIFIESTOS, "*.xlsx"))
        if archivos_excel:
            return max(archivos_excel, key=os.path.getctime)
            
    return ruta_manifiesto

def obtener_json_puerto():
    ruta_manifiesto = obtener_ruta_manifiesto()
    
    if not os.path.exists(ruta_manifiesto):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Codigos_Esperados", "Naviera", "Estado"])
        wb.save(ruta_manifiesto)

    manifiesto_json = {}
    try:
        wb = openpyxl.load_workbook(ruta_manifiesto, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                codigo = str(row[0]).strip()
                manifiesto_json[codigo] = {
                    "naviera": str(row[1]).strip() if len(row) > 1 and row[1] else "GENERICO",
                    "estado": str(row[2]).strip() if len(row) > 2 and row[2] else "Pendiente"
                }
    except Exception as e:
        print(f"❌ Error leyendo manifiesto: {e}")
    return manifiesto_json

def actualizar_estado_puerto(codigo_ocr, nuevo_estado):
    ruta_manifiesto = obtener_ruta_manifiesto()
    try:
        wb = openpyxl.load_workbook(ruta_manifiesto)
        ws = wb.active
        codigo_limpio_ocr = codigo_ocr.replace("-", "").replace(" ", "").upper()
        for row in ws.iter_rows(min_row=2):
            if row[0].value and str(row[0].value).replace("-", "").replace(" ", "").upper() == codigo_limpio_ocr:
                row[2].value = nuevo_estado
                break
        wb.save(ruta_manifiesto)
    except Exception as e:
        print(f"❌ Error al actualizar estado: {e}")

def registrar_acceso_denegado(datos_precinto, motivo):
    ruta_manifiesto = obtener_ruta_manifiesto()
    
    # 1. Si no existe, creamos el archivo desde cero
    if not os.path.exists(ruta_manifiesto):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Manifiesto"
        ws.append(["Codigos_Esperados", "Naviera", "Estado"])
    else:
        # 2. Si existe, lo cargamos
        try:
            wb = openpyxl.load_workbook(ruta_manifiesto)
        except Exception as e:
            print(f"❌ No se pudo abrir el Excel: {e}")
            return
    
    # 3. Asegurar la hoja "Denegados"
    if "Denegados" not in wb.sheetnames:
        ws2 = wb.create_sheet(title="Denegados")
        ws2.append(["Fecha y Hora", "Código Detectado", "Naviera", "Confianza", "Motivo"])
    else:
        ws2 = wb["Denegados"]
        
    ws2.append([
        datos_precinto["timestamp"],
        datos_precinto.get("codigo_precinto") or "NO LEGIBLE",
        datos_precinto.get("naviera") or "GENERICO",
        datos_precinto["confianza_lectura"],
        motivo
    ])
    
    # 4. Guardado seguro
    try:
        wb.save(ruta_manifiesto)
        print(f"✅ Registro guardado en Hoja 2.")
    except Exception as e:
        print(f"❌ Error al guardar (¿Tienes el Excel abierto?): {e}")
# ==============================================================================
# LÓGICA OCR (PROCESAMIENTO EN SEGUNDO PLANO)
# ==============================================================================

resultado_validacion_texto = ""
resultado_validacion_color = (255, 255, 255)

def procesar_ocr_bg(ruta):
    global resultado_validacion_texto, resultado_validacion_color
    img = cv2.imread(ruta)
    if img is None: return

    texto_crudo, score_val, variante = mejor_ocr(img)
    datos_precinto = {
        "timestamp": datetime.now().isoformat(),
        "codigo_precinto": None, "naviera": None,
        "confianza_lectura": round(score_val, 2),
        "variante_imagen": variante,
        "texto_crudo_completo": texto_crudo.replace('\n', ' | ') 
    }

    # Extracción y limpieza
    patron_codigo = r'([A-Z\-\.\+]{1,9}\s*[0-9]{5,10})'
    for linea in texto_crudo.split('\n'):
        linea_limpia = linea.strip()
        if "MAERSK" in linea_limpia.upper(): datos_precinto["naviera"] = "MAERSK"
        elif "MSC" in linea_limpia.upper(): datos_precinto["naviera"] = "MSC"
        elif "EVONIK" in linea_limpia.upper(): datos_precinto["naviera"] = "EVONIK"
        
        coincidencia = re.search(patron_codigo, linea_limpia)
        if coincidencia:
             codigo_encontrado = coincidencia.group(1).replace(" ", "").replace(".", "").replace("+", "-")
             if len(codigo_encontrado) > 5 and any(c.isdigit() for c in codigo_encontrado):
                 if datos_precinto["codigo_precinto"] is None: datos_precinto["codigo_precinto"] = codigo_encontrado

    # Validación
    codigo_detectado = datos_precinto.get("codigo_precinto")
    if datos_precinto["naviera"] is None: datos_precinto["naviera"] = "GENERICO"

    if codigo_detectado:
        json_puerto = obtener_json_puerto()
        codigo_limpio = codigo_detectado.replace("-", "").replace(" ", "").upper()
        esperados = [str(c).replace("-", "").replace(" ", "").upper() for c in json_puerto.keys()]
        
        if codigo_limpio in esperados:
            resultado_validacion_texto = f"OK: {codigo_detectado} - MATCH"
            resultado_validacion_color = (0, 255, 0)
            actualizar_estado_puerto(codigo_detectado, "Camión Pasó")
        else:
            resultado_validacion_texto = f"ERROR: {codigo_detectado} - NO DECLARADO"
            resultado_validacion_color = (0, 0, 255)
            registrar_acceso_denegado(datos_precinto, "No en manifiesto")
    else:
        resultado_validacion_texto = "ERROR: No legible"
        resultado_validacion_color = (0, 165, 255)
        registrar_acceso_denegado(datos_precinto, "Fallo de lectura OCR")


def ejecutar_vision():
    global resultado_validacion_texto, resultado_validacion_color
    
    user = os.getenv("CAM_USER")
    pas = os.getenv("CAM_PASS")
    ip = os.getenv("CAM_IP")
    
    service = CameraService()
    rtsp_url = f"rtsp://{user}:{pas}@{ip}:554/Streaming/Channels/101"
    
    service.configurar_auto_iris(ip)
    threading.Thread(target=service.lector_rtsp, args=(rtsp_url,), daemon=True).start()

    vel_mov = 30
    vel_lente = 5
    estado_actual = "CONECTADO (IRIS AUTO)"
    
    WINDOW_NAME = "ACCELERA SVTI - Control PTZ"
    cv2.namedWindow(WINDOW_NAME, cv2.WINDOW_NORMAL)
    cv2.resizeWindow(WINDOW_NAME, 1280, 720)

    while True:
        with service.lock:
            frame = None if service.ultimo_frame is None else service.ultimo_frame.copy()

        if frame is None:
            espera = np.zeros((720, 1280, 3), dtype=np.uint8)
            cv2.putText(espera, "Conectando...", (400, 360), cv2.FONT_HERSHEY_SIMPLEX, 1, (0, 255, 255), 2)
            cv2.imshow(WINDOW_NAME, espera)
            if cv2.waitKey(1) & 0xFF == 27: break
            continue

        frame_limpio = frame.copy()
        rect = cv2.getWindowImageRect(WINDOW_NAME)
        ancho_v, alto_v = max(800, rect[2]), max(450, rect[3])
        alto_video = alto_v - 110
        frame_achicado = cv2.resize(frame, (ancho_v, alto_video))
        
        centro_x, centro_y_s = int(ancho_v/2), int(alto_video/2)
        ancho_r, alto_r = int(ancho_v * ANCHO_RETICULA_PCT), int(alto_video * ALTO_RETICULA_PCT)
        y1, y2 = max(0, centro_y_s - int(alto_r/2)), min(alto_video, centro_y_s + int(alto_r/2))
        x1, x2 = max(0, centro_x - int(ancho_r/2)), min(ancho_v, centro_x + int(ancho_r/2))
        
        frame_mostrar = cv2.copyMakeBorder(frame_achicado, 50, 60, 0, 0, cv2.BORDER_CONSTANT, value=[15, 15, 25])
        cv2.rectangle(frame_mostrar, (x1, y1 + 50), (x2, y2 + 50), COLOR_RETICULA, GROSOR_RETICULA)
        
        cv2.putText(frame_mostrar, f"ESTADO: {estado_actual}", (20, 30), cv2.FONT_HERSHEY_SIMPLEX, 0.5, (255, 255, 255), 1)
        cv2.putText(frame_mostrar, f"VEL: {vel_mov}", (250, 30), cv2.FONT_HERSHEY_SIMPLEX, 0.5, (0, 200, 255), 1)
        cv2.putText(frame_mostrar, "MOV: [8456] | ZOOM: [WS] | FOCO: [79] | PARAR: [ESPACIO]", (20, frame_mostrar.shape[0]-35), cv2.FONT_HERSHEY_SIMPLEX, 0.45, (200, 200, 200), 1)
        cv2.putText(frame_mostrar, "SISTEMA: OCR [P] | VEL [1-3] | SALIR [ESC]", (20, frame_mostrar.shape[0]-15), cv2.FONT_HERSHEY_SIMPLEX, 0.45, (200, 200, 200), 1)

        if resultado_validacion_texto != "":
            tamano_texto, _ = cv2.getTextSize(resultado_validacion_texto, cv2.FONT_HERSHEY_SIMPLEX, 0.8, 2)
            cv2.rectangle(frame_mostrar, (45, alto_video - 45), (55 + tamano_texto[0], alto_video - 15), (0, 0, 0), -1)
            cv2.putText(frame_mostrar, resultado_validacion_texto, (50, alto_video - 20), cv2.FONT_HERSHEY_SIMPLEX, 0.8, resultado_validacion_color, 2)

        cv2.imshow(WINDOW_NAME, frame_mostrar)
        tecla = cv2.waitKey(1) & 0xFF

        if tecla == 27: break
        elif tecla == 32: 
            service.mover_camara(ip, 0, 0, 0, 0, 0)
            estado_actual = "DETENIDO"
        elif tecla == ord('1'): vel_mov = 20
        elif tecla == ord('2'): vel_mov = 50
        elif tecla == ord('3'): vel_mov = 100
        elif tecla == ord('8'): 
            service.mover_camara(ip, pan=0, tilt=vel_mov)
            estado_actual = "SUBIENDO"
        elif tecla == ord('5'): 
            service.mover_camara(ip, pan=0, tilt=-vel_mov)
            estado_actual = "BAJANDO"
        elif tecla == ord('4'): 
            service.mover_camara(ip, pan=-vel_mov, tilt=0)
            estado_actual = "IZQUIERDA"
        elif tecla == ord('6'): 
            service.mover_camara(ip, pan=vel_mov, tilt=0)
            estado_actual = "DERECHA"
        elif tecla == ord('w') or tecla == ord('W'): 
            service.mover_camara(ip, zoom=vel_mov)
            estado_actual = "ZOOM IN"
        elif tecla == ord('s') or tecla == ord('S'): 
            service.mover_camara(ip, zoom=-vel_mov)
            estado_actual = "ZOOM OUT"
        elif tecla == ord('7'): 
            service.mover_camara(ip, focus=vel_lente)
            estado_actual = "FOCO CERCA (+)"
        elif tecla == ord('9'): 
            service.mover_camara(ip, focus=-vel_lente)
            estado_actual = "FOCO LEJOS (-)"
        elif tecla == ord('p') or tecla == ord('P'):
            resultado_validacion_texto = "Procesando validacion..."
            resultado_validacion_color = (0, 255, 255)
            ruta = os.path.join(CARPETA_CAPTURAS, f"precinto_{datetime.now().strftime('%Y%m%d_%H%M%S')}.jpg")
            cv2.imwrite(ruta, frame_limpio)
            threading.Thread(target=procesar_ocr_bg, args=(ruta,), daemon=True).start()
            estado_actual = "CAPTURANDO OCR..."

    service.mover_camara(ip, 0, 0, 0, 0, 0)
    cv2.destroyAllWindows()